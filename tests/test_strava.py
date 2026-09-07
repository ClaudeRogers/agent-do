#!/usr/bin/env python3
import importlib.machinery, importlib.util, json, os, subprocess, tempfile, threading
from pathlib import Path
from urllib.request import urlopen

ROOT = Path(__file__).resolve().parents[1]
TOOL = ROOT / "tools" / "agent-strava"
strava = importlib.machinery.SourceFileLoader("agent_strava", str(TOOL)).load_module()

def run(*args, env=None):
    return subprocess.run([str(TOOL), *args], text=True, capture_output=True, env=env)

def test_status_without_profile():
    with tempfile.TemporaryDirectory() as home:
        result = run("status", env={**os.environ, "AGENT_DO_HOME": home})
        assert result.returncode == 0 and "No local Strava profile" in result.stdout

def test_status_json_without_profile_is_machine_readable():
    with tempfile.TemporaryDirectory() as home:
        result = run("status", "--json", env={**os.environ, "AGENT_DO_HOME": home})
        assert result.returncode == 0
        assert json.loads(result.stdout) == {"configured": False, "cache_present": False, "next_command": "agent-do strava init"}

def test_profile_units_are_local_configuration():
    with tempfile.TemporaryDirectory() as home:
        data = Path(home) / "strava"; data.mkdir()
        (data / "profile.json").write_text(json.dumps({"athlete_name": "Test", "units": "metric"}))
        result = run("status", "--json", env={**os.environ, "AGENT_DO_HOME": home})
        assert result.returncode == 0 and json.loads(result.stdout)["units"] == "metric"

def test_dashboard_uses_local_cache_only():
    with tempfile.TemporaryDirectory() as home:
        data = Path(home) / "strava"; data.mkdir()
        (data / "activities.json").write_text(json.dumps({"synced_at": "2026-09-06T00:00:00+00:00", "activities": []}))
        result = run("dashboard", env={**os.environ, "AGENT_DO_HOME": home})
        assert result.returncode == 0
        page = Path(result.stdout.strip())
        assert page.exists() and "Your training" in page.read_text()

def test_export_csv_uses_local_cache_and_omits_sensitive_route_fields():
    with tempfile.TemporaryDirectory() as home:
        data = Path(home) / "strava"; data.mkdir()
        (data / "profile.json").write_text(json.dumps({"units": "metric"}))
        now = strava.datetime.now(strava.timezone.utc).isoformat()
        (data / "activities.json").write_text(json.dumps({"synced_at": now, "activities": [{"id": 42, "name": "Home address run", "start_date": now, "sport_type": "Run", "distance": 5000, "moving_time": 1800, "map": {"summary_polyline": "secret-route"}}]}))
        output = Path(home) / "activities.csv"
        result = run("export", str(output), env={**os.environ, "AGENT_DO_HOME": home})
        assert result.returncode == 0 and json.loads(result.stdout)["activities"] == 1
        content = output.read_text()
        assert "Distance (km)" in content and "Moving time" in content and "0:30:00" in content
        assert "secret-route" not in content and "Home address" not in content
        duplicate = run("export", str(output), env={**os.environ, "AGENT_DO_HOME": home})
        assert duplicate.returncode != 0 and "--overwrite" in duplicate.stderr

def test_export_activity_filter_accepts_comma_and_bracketed_groups():
    with tempfile.TemporaryDirectory() as home:
        data = Path(home) / "strava"; data.mkdir()
        (data / "profile.json").write_text(json.dumps({"units": "metric"}))
        now = strava.datetime.now(strava.timezone.utc).isoformat()
        activities = [
            {"id": 1, "start_date": now, "sport_type": "TrailRun", "distance": 5000, "moving_time": 1800},
            {"id": 2, "start_date": now, "sport_type": "Ride", "distance": 20000, "moving_time": 3600},
            {"id": 3, "start_date": now, "sport_type": "WeightTraining", "moving_time": 1200},
        ]
        (data / "activities.json").write_text(json.dumps({"synced_at": now, "activities": activities}))
        output = Path(home) / "run-bike.csv"
        result = run("export", str(output), "--activity", "[run,bike]", env={**os.environ, "AGENT_DO_HOME": home})
        assert result.returncode == 0 and json.loads(result.stdout)["activities"] == 2
        content = output.read_text()
        assert "TrailRun" in content and "Ride" in content and "WeightTraining" not in content

def test_export_xlsx_has_readable_summary_and_activity_sheets():
    if importlib.util.find_spec("openpyxl") is None:
        return
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as home:
        data = Path(home) / "strava"; data.mkdir()
        (data / "profile.json").write_text(json.dumps({"units": "imperial"}))
        now = strava.datetime.now(strava.timezone.utc).isoformat()
        (data / "activities.json").write_text(json.dumps({"synced_at": now, "activities": [{"id": 42, "start_date": now, "sport_type": "Ride", "distance": 16093, "moving_time": 3600}]}))
        output = Path(home) / "activities.xlsx"
        result = run("export", str(output), env={**os.environ, "AGENT_DO_HOME": home})
        assert result.returncode == 0
        workbook = load_workbook(output)
        assert workbook.sheetnames == ["Summary", "Activities", "Weekly", "Monthly", "Data dictionary"]
        assert workbook["Summary"]["A8"].value == "Distance (mi)"
        assert workbook["Weekly"]["B1"].value == "Distance (mi)"
        assert workbook["Monthly"]["E1"].value == "Elevation gain (ft)"
        assert workbook["Activities"]["D1"].value == "Moving time"
        assert workbook["Activities"]["D2"].value == "01:00:00"
        assert workbook["Summary"]["B9"].value == "01:00:00"
        assert workbook["Activities"].max_row == 2

def test_gear_export_xlsx_has_summary_then_one_activity_sheet_per_gear():
    if importlib.util.find_spec("openpyxl") is None:
        return
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as home:
        data = Path(home) / "strava"; data.mkdir()
        (data / "profile.json").write_text(json.dumps({"units": "imperial"}))
        now = strava.datetime.now(strava.timezone.utc).isoformat()
        activities = [
            {"id": 1, "start_date": now, "sport_type": "Run", "distance": 5000, "moving_time": 1800, "kilojoules": 418.4, "gear_id": "g1"},
            {"id": 2, "start_date": now, "sport_type": "Run", "distance": 3000, "moving_time": 1200, "gear_id": "g2"},
        ]
        gear = [
            {"id": "g1", "name": "Road shoes", "distance": 80467},
            {"id": "g2", "name": "Trail shoes", "distance": 32187},
        ]
        (data / "activities.json").write_text(json.dumps({"synced_at": now, "gear": gear, "activities": activities}))
        output = Path(home) / "gear.xlsx"
        result = run("gear-export", str(output), env={**os.environ, "AGENT_DO_HOME": home})
        assert result.returncode == 0
        assert json.loads(result.stdout) == {"exported": str(output), "format": "xlsx", "gear": 2, "activities": 2, "units": "imperial"}
        workbook = load_workbook(output)
        assert workbook.sheetnames == ["Gear", "Road shoes", "Trail shoes"]
        assert [cell.value for cell in workbook["Gear"][1]] == ["Gear", "Total distance (mi)", "Last used", "Activities"]
        assert workbook["Gear"]["A2"].value == "Road shoes"
        assert workbook["Gear"]["B2"].value == 50.0
        assert workbook["Gear"]["D3"].value == 1
        assert workbook["Road shoes"]["A1"].value == "Date"
        assert workbook["Road shoes"].max_row == 2
        headers = [cell.value for cell in workbook["Road shoes"][1]]
        assert workbook["Road shoes"].cell(2, headers.index("Moving time") + 1).value == "30:00"
        assert workbook["Road shoes"].cell(2, headers.index("Average pace (min:sec/mi)") + 1).value == "09:39"
        assert workbook["Road shoes"].cell(2, headers.index("Calories") + 1).value == 100

def test_export_rows_estimates_calories_from_strava_kilojoules():
    rows = strava.export_rows([{"id": 1, "start_date": "2026-09-07T00:00:00+00:00", "sport_type": "Run", "distance": 5000, "moving_time": 1800, "kilojoules": 418.4}], "imperial")
    assert rows[0]["Calories"] == 100
    assert strava.format_export_duration(3599) == "59:59"
    assert strava.format_export_duration(3600) == "01:00:00"

def test_dashboard_export_uses_selected_local_cache_without_retaining_a_file():
    with tempfile.TemporaryDirectory() as home:
        original = strava.HOME, strava.PROFILE, strava.CACHE
        try:
            strava.HOME = Path(home) / "strava"; strava.HOME.mkdir()
            strava.PROFILE, strava.CACHE = strava.HOME / "profile.json", strava.HOME / "activities.json"
            strava.PROFILE.write_text(json.dumps({"units": "imperial"}))
            now = strava.datetime.now(strava.timezone.utc).isoformat()
            strava.CACHE.write_text(json.dumps({"synced_at": now, "activities": [{"id": 42, "start_date": now, "sport_type": "Run", "distance": 5000, "moving_time": 1800}]}))
            content, count = strava.export_download("csv", 30, "Run")
            assert count == 1 and b"0:30:00" in content and b"summary_polyline" not in content
        finally:
            strava.HOME, strava.PROFILE, strava.CACHE = original

def test_dashboard_export_endpoint_downloads_the_selected_csv():
    with tempfile.TemporaryDirectory() as home:
        original = strava.HOME, strava.PROFILE, strava.CACHE
        server = None
        try:
            strava.HOME = Path(home) / "strava"; strava.HOME.mkdir()
            strava.PROFILE, strava.CACHE = strava.HOME / "profile.json", strava.HOME / "activities.json"
            strava.PROFILE.write_text(json.dumps({"units": "metric"}))
            now = strava.datetime.now(strava.timezone.utc).isoformat()
            strava.CACHE.write_text(json.dumps({"synced_at": now, "activities": [{"id": 42, "start_date": now, "sport_type": "Run", "distance": 5000, "moving_time": 1800}]}))
            server = strava.HTTPServer(("127.0.0.1", 0), strava.LocalDashboard)
            worker = threading.Thread(target=server.handle_request); worker.start()
            with urlopen(f"http://127.0.0.1:{server.server_port}/api/export?format=csv&days=30&type=Run", timeout=5) as response:
                assert response.status == 200 and "attachment" in response.headers["Content-Disposition"]
                assert b"0:30:00" in response.read()
            worker.join(timeout=5)
        finally:
            if server: server.server_close()
            strava.HOME, strava.PROFILE, strava.CACHE = original

def test_summary_calculates_selected_range():
    now = strava.datetime.now(strava.timezone.utc)
    activities = [
        {"start_date": (now - strava.timedelta(days=2)).isoformat(), "distance": 5000, "moving_time": 1800, "total_elevation_gain": 100},
        {"start_date": (now - strava.timedelta(days=10)).isoformat(), "distance": 9000, "moving_time": 3600, "total_elevation_gain": 200},
    ]
    result = strava.summarize(activities, days=7)
    assert result["distance_km"] == 5 and result["moving_minutes"] == 30 and result["activity_count"] == 1
    assert result["weekly"] and result["weekly"][0]["distance_km"] == 5

def test_summary_filters_by_specific_strava_sport_type():
    now = strava.datetime.now(strava.timezone.utc)
    activities = [
        {"start_date": (now - strava.timedelta(days=1)).isoformat(), "sport_type": "TrailRun", "type": "Run", "distance": 5000},
        {"start_date": (now - strava.timedelta(days=1)).isoformat(), "sport_type": "Ride", "type": "Ride", "distance": 12000},
    ]
    result = strava.summarize(activities, days=7, activity_type="TrailRun")
    assert result["activity_count"] == 1 and result["distance_km"] == 5
    assert strava.activity_kind({"sport_type": "WeightTraining"}) == "WeightTraining"

def test_summary_pace_uses_only_run_and_walk_activities():
    now = strava.datetime.now(strava.timezone.utc)
    activities = [
        {"start_date": (now - strava.timedelta(days=1)).isoformat(), "sport_type": "Run", "distance": 5000, "moving_time": 1800},
        {"start_date": (now - strava.timedelta(days=1)).isoformat(), "sport_type": "Walk", "distance": 1000, "moving_time": 600},
        {"start_date": (now - strava.timedelta(days=1)).isoformat(), "sport_type": "Ride", "distance": 20000, "moving_time": 3600},
    ]
    result = strava.summarize(activities, days=7)
    assert result["pace_metric"] == "pace"
    assert result["weekly"][0]["pace_seconds_per_km"] == 400
    assert result["pace_seconds_per_km"] == 400
    bike = strava.summarize(activities, days=7, activity_type="Ride")
    assert bike["pace_metric"] == "speed" and bike["weekly"][0]["average_speed_mps"] == 5.56
    assert strava.summarize(activities, days=7, activity_type="Swim")["pace_metric"] is None

def test_gear_summary_uses_strava_lifetime_distance_and_cached_activity_history():
    now = strava.datetime.now(strava.timezone.utc).isoformat()
    cache = {"gear": [{"id": "g1", "name": "Road shoes", "distance": 80467}], "activities": [
        {"id": 1, "start_date": now, "distance": 5000, "gear": {"id": "g1", "name": "Road shoes"}},
        {"id": 2, "start_date": now, "distance": 3000, "gear": {"id": "g1", "name": "Road shoes"}},
    ]}
    gear = strava.gear_summary(cache)
    assert gear == [{"id": "g1", "name": "Road shoes", "brand_name": None, "model_name": None, "primary": False, "distance_m": 80467, "activity_count": 2, "last_used_at": now}]
    assert [activity["id"] for activity in strava.activities_for_gear(cache, "g1")] == [1, 2]
    assert strava.gear_id({"gear_id": "g1"}) == "g1"

def test_gear_api_returns_lifetime_distance_and_associated_cached_activities():
    with tempfile.TemporaryDirectory() as home:
        original = strava.HOME, strava.PROFILE, strava.CACHE
        server = None
        try:
            strava.HOME = Path(home) / "strava"; strava.HOME.mkdir()
            strava.PROFILE, strava.CACHE = strava.HOME / "profile.json", strava.HOME / "activities.json"
            strava.PROFILE.write_text(json.dumps({"units": "imperial"}))
            now = strava.datetime.now(strava.timezone.utc).isoformat()
            strava.CACHE.write_text(json.dumps({"synced_at": now, "gear": [{"id": "g1", "name": "Road shoes", "distance": 80467}], "activities": [{"id": 42, "start_date": now, "distance": 5000, "gear": {"id": "g1", "name": "Road shoes"}}]}))
            server = strava.HTTPServer(("127.0.0.1", 0), strava.LocalDashboard)
            worker = threading.Thread(target=server.handle_request); worker.start()
            with urlopen(f"http://127.0.0.1:{server.server_port}/api/gear", timeout=5) as response:
                data = json.load(response)
            worker.join(timeout=5)
            assert data["units"] == "imperial" and data["gear"][0]["distance_m"] == 80467
            worker = threading.Thread(target=server.handle_request); worker.start()
            with urlopen(f"http://127.0.0.1:{server.server_port}/api/gear/g1/activities", timeout=5) as response:
                data = json.load(response)
            worker.join(timeout=5)
            assert [activity["id"] for activity in data["activities"]] == [42]
        finally:
            if server: server.server_close()
            strava.HOME, strava.PROFILE, strava.CACHE = original

def test_gear_export_endpoint_downloads_an_excel_workbook():
    with tempfile.TemporaryDirectory() as home:
        original = strava.HOME, strava.PROFILE, strava.CACHE
        server = None
        try:
            strava.HOME = Path(home) / "strava"; strava.HOME.mkdir()
            strava.PROFILE, strava.CACHE = strava.HOME / "profile.json", strava.HOME / "activities.json"
            strava.PROFILE.write_text(json.dumps({"units": "imperial"}))
            now = strava.datetime.now(strava.timezone.utc).isoformat()
            strava.CACHE.write_text(json.dumps({"synced_at": now, "gear": [{"id": "g1", "name": "Road shoes", "distance": 80467}], "activities": [{"id": 42, "start_date": now, "sport_type": "Run", "distance": 5000, "moving_time": 1800, "gear_id": "g1"}]}))
            server = strava.HTTPServer(("127.0.0.1", 0), strava.LocalDashboard)
            worker = threading.Thread(target=server.handle_request); worker.start()
            with urlopen(f"http://127.0.0.1:{server.server_port}/api/gear-export", timeout=5) as response:
                content = response.read()
                assert response.status == 200 and "attachment" in response.headers["Content-Disposition"]
            worker.join(timeout=5)
            assert content.startswith(b"PK")
        finally:
            if server: server.server_close()
            strava.HOME, strava.PROFILE, strava.CACHE = original

def test_responsive_dashboard_uses_manual_sync_without_polling():
    assert strava.DYNAMIC_DASHBOARD_PATH.name == "agent-strava-dashboard.html"
    assert strava.DYNAMIC_DASHBOARD_PATH.parent == ROOT / "docs"
    assert "/api/sync" in strava.DYNAMIC_DASHBOARD
    assert "/api/export" in strava.DYNAMIC_DASHBOARD
    assert "Export Excel" in strava.DYNAMIC_DASHBOARD
    assert "Export CSV" in strava.DYNAMIC_DASHBOARD
    assert "/api/preferences" in strava.DYNAMIC_DASHBOARD
    assert "activity-type" in strava.DYNAMIC_DASHBOARD
    assert "activity_types" in strava.DYNAMIC_DASHBOARD
    assert "activity-heading" in strava.DYNAMIC_DASHBOARD
    assert "friendly" in strava.DYNAMIC_DASHBOARD
    assert "Date &amp; time" in strava.DYNAMIC_DASHBOARD
    assert "WeightTraining'?'--'" in strava.DYNAMIC_DASHBOARD
    assert "setInterval" not in strava.DYNAMIC_DASHBOARD
    assert "localStorage" not in strava.DYNAMIC_DASHBOARD
    assert "Units: Imperial" in strava.DYNAMIC_DASHBOARD
    assert '<button data-days=7>1 week</button>' in strava.DYNAMIC_DASHBOARD
    assert '<button class=active data-days=30>1 month</button>' in strava.DYNAMIC_DASHBOARD
    assert '<button data-days=90>3 months</button>' in strava.DYNAMIC_DASHBOARD
    assert '<button data-days=365>1 year</button>' in strava.DYNAMIC_DASHBOARD
    assert "let days=30," in strava.DYNAMIC_DASHBOARD
    assert "4 weeks" not in strava.DYNAMIC_DASHBOARD
    assert "activity-pagination" in strava.DYNAMIC_DASHBOARD
    assert "ACTIVITIES_PER_PAGE=10" in strava.DYNAMIC_DASHBOARD
    assert "Showing '+(first+1)+'–'+last+' of '+activities.length" in strava.DYNAMIC_DASHBOARD
    assert "</table><nav id=activity-pagination" in strava.DYNAMIC_DASHBOARD
    assert "function weekLabelStep(count)" in strava.DYNAMIC_DASHBOARD
    assert "count<=26?4:8" in strava.DYNAMIC_DASHBOARD
    assert "index%labelStep===0||index===weeks.length-1" in strava.DYNAMIC_DASHBOARD
    assert "week-pace-card" in strava.DYNAMIC_DASHBOARD
    assert "weekly-charts" in strava.DYNAMIC_DASHBOARD
    assert "function updateOneWeekSummary(data)" in strava.DYNAMIC_DASHBOARD
    assert "paceChart.hidden=true" in strava.DYNAMIC_DASHBOARD
    assert "charts.style.display=oneWeek?'none':''" in strava.DYNAMIC_DASHBOARD
    assert "Distance per week" in strava.DYNAMIC_DASHBOARD
    assert "Moving time per week" in strava.DYNAMIC_DASHBOARD
    assert "Average pace per week" in strava.DYNAMIC_DASHBOARD
    assert "pace-card" in strava.DYNAMIC_DASHBOARD
    assert "pace-axis" in strava.DYNAMIC_DASHBOARD
    assert "distance-axis" in strava.DYNAMIC_DASHBOARD
    assert "time-axis" in strava.DYNAMIC_DASHBOARD
    assert "drawAxis" in strava.DYNAMIC_DASHBOARD
    assert "niceStep" in strava.DYNAMIC_DASHBOARD
    assert "chart-grid" in strava.DYNAMIC_DASHBOARD
    assert "chart-tooltip" in strava.DYNAMIC_DASHBOARD
    assert "plotHeight=147" in strava.DYNAMIC_DASHBOARD
    assert "height:147px" in strava.DYNAMIC_DASHBOARD
    assert "align-items:center;justify-content:flex-end" in strava.DYNAMIC_DASHBOARD
    assert "scale.values[index]/scale.max*147)+'px'" in strava.DYNAMIC_DASHBOARD
    assert "weekLabel(week.week)" in strava.DYNAMIC_DASHBOARD
    assert "align-items:flex-start" in strava.DYNAMIC_DASHBOARD
    assert "localRoute" in strava.DYNAMIC_DASHBOARD
    assert "Route map" in strava.DYNAMIC_DASHBOARD
    assert "maplibre-gl@5" in strava.DYNAMIC_DASHBOARD
    assert "tiles.openfreemap.org/styles/bright" in strava.DYNAMIC_DASHBOARD
    assert "activity-route-casing" in strava.DYNAMIC_DASHBOARD
    assert "activityMap.fitBounds" in strava.DYNAMIC_DASHBOARD
    assert "routeFallback" in strava.DYNAMIC_DASHBOARD
    assert "leaflet@1.9.4" not in strava.DYNAMIC_DASHBOARD
    assert "Interactive map © OpenFreeMap" in strava.DYNAMIC_DASHBOARD
    assert "CARTO" not in strava.DYNAMIC_DASHBOARD
    assert "Splits" in strava.DYNAMIC_DASHBOARD
    assert "Best efforts" in strava.DYNAMIC_DASHBOARD
    assert "velocity_smooth" in strava.DYNAMIC_DASHBOARD
    assert strava.DYNAMIC_DASHBOARD.index("Recent activity") < strava.DYNAMIC_DASHBOARD.index("Distance per week")
    assert "column-grip" in strava.DYNAMIC_DASHBOARD
    assert "th+th{border-left" in strava.DYNAMIC_DASHBOARD
    assert "td+td,th+th" not in strava.DYNAMIC_DASHBOARD
    assert "text-transform:uppercase" in strava.DYNAMIC_DASHBOARD
    assert "hour'+(hours===1?'':'s')" in strava.DYNAMIC_DASHBOARD
    assert "activity-dialog" in strava.DYNAMIC_DASHBOARD
    assert "/api/activity/" in strava.DYNAMIC_DASHBOARD
    assert "'/gear'" in strava.DYNAMIC_DASHBOARD
    assert "/api/gear" in strava.DYNAMIC_DASHBOARD
    assert "Gear activity history" in strava.DYNAMIC_DASHBOARD
    assert "#export-xlsx,#export-csv" in strava.DYNAMIC_DASHBOARD
    assert "gear-export" in strava.DYNAMIC_DASHBOARD
    assert "Export Gear" in strava.DYNAMIC_DASHBOARD
    assert "dialog.dataset.backdropClose" in strava.DYNAMIC_DASHBOARD
    assert "dialog.getBoundingClientRect()" in strava.DYNAMIC_DASHBOARD
    assert "document.addEventListener('pointerdown'" in strava.DYNAMIC_DASHBOARD
    assert "stream-axis" in strava.DYNAMIC_DASHBOARD
    assert "line.setAttribute('stroke-linecap','round')" in strava.DYNAMIC_DASHBOARD
    assert "const smoothed=values.map" in strava.DYNAMIC_DASHBOARD
    assert "cell.title=cell.textContent" in strava.DYNAMIC_DASHBOARD
    assert "white-space:nowrap" in strava.DYNAMIC_DASHBOARD
    assert "font-size:clamp(1.55rem,2.5vw,2.1rem)" in strava.DYNAMIC_DASHBOARD
    assert "Number(value).toLocaleString" in strava.DYNAMIC_DASHBOARD
    assert "number(Math.round(m*3.28084))+' ft'" in strava.DYNAMIC_DASHBOARD
    assert "resolvedOptions().timeZone" in strava.DYNAMIC_DASHBOARD
    assert "new Date(a.start_date).toLocaleString" in strava.DYNAMIC_DASHBOARD
    assert "low=Math.max(0,minimum-padding)" in strava.DYNAMIC_DASHBOARD
    assert "filter(([,value])=>present(value))" in strava.DYNAMIC_DASHBOARD
    result = run("serve", "--help")
    assert result.returncode == 0 and "--no-sync" in result.stdout

def test_connect_requests_private_activity_scope():
    assert "activity:read,activity:read_all" in strava.connect.__code__.co_consts

if __name__ == "__main__":
    test_status_without_profile(); test_status_json_without_profile_is_machine_readable(); test_profile_units_are_local_configuration(); test_dashboard_uses_local_cache_only(); test_export_csv_uses_local_cache_and_omits_sensitive_route_fields(); test_export_activity_filter_accepts_comma_and_bracketed_groups(); test_export_xlsx_has_readable_summary_and_activity_sheets(); test_gear_export_xlsx_has_summary_then_one_activity_sheet_per_gear(); test_export_rows_estimates_calories_from_strava_kilojoules(); test_dashboard_export_uses_selected_local_cache_without_retaining_a_file(); test_dashboard_export_endpoint_downloads_the_selected_csv(); test_summary_calculates_selected_range(); test_summary_filters_by_specific_strava_sport_type(); test_summary_pace_uses_only_run_and_walk_activities(); test_gear_summary_uses_strava_lifetime_distance_and_cached_activity_history(); test_gear_api_returns_lifetime_distance_and_associated_cached_activities(); test_gear_export_endpoint_downloads_an_excel_workbook(); test_responsive_dashboard_uses_manual_sync_without_polling(); test_connect_requests_private_activity_scope(); print("ok")
